from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()

HEADER_FILL = PatternFill('solid', fgColor='1E1B4B')
HEADER_FONT = Font(name='Arial', bold=True, color='FFFFFF', size=11)
TITLE_FONT = Font(name='Arial', bold=True, size=14, color='1E1B4B')
SECTION_FONT = Font(name='Arial', bold=True, size=12, color='4F46E5')
BOLD = Font(name='Arial', bold=True, size=10)
NORMAL = Font(name='Arial', size=10)
FORMULA_FONT = Font(name='Consolas', size=11, color='0000FF')
LIGHT_FILL = PatternFill('solid', fgColor='F0F0FF')
ORANGE_FILL = PatternFill('solid', fgColor='FFF7ED')
GREEN_FILL = PatternFill('solid', fgColor='F0FFF0')
YELLOW_FILL = PatternFill('solid', fgColor='FFFFF0')
thin = Side(style='thin', color='CCCCCC')
BORDER = Border(top=thin, bottom=thin, left=thin, right=thin)

def write_row(ws, row, data, font=NORMAL, fill=None):
    for i, v in enumerate(data, 1):
        cell = ws.cell(row=row, column=i, value=v)
        cell.font = font
        if fill:
            cell.fill = fill
        cell.border = BORDER
        cell.alignment = Alignment(vertical='center', wrap_text=True)

# ===== Sheet 1 =====
ws1 = wb.active
ws1.title = '計算ルール概要'
ws1.sheet_properties.tabColor = '4F46E5'
ws1.column_dimensions['A'].width = 30
ws1.column_dimensions['B'].width = 50
ws1.column_dimensions['C'].width = 20

ws1.cell(row=1, column=1, value='VAVITTE KB計算ルール 最終確定版').font = TITLE_FONT
ws1.cell(row=2, column=1, value='作成日: 2026年4月13日').font = Font(name='Arial', size=9, color='666666')

ws1.cell(row=4, column=1, value='基本計算式').font = SECTION_FONT
ws1.cell(row=5, column=1, value='KB = 定価 x (サロン価格率 - 代理店原価率) x 数量').font = FORMULA_FONT

ws1.cell(row=7, column=1, value='定価の逆算').font = SECTION_FONT
ws1.cell(row=8, column=1, value='定価 = BカートAPI最高単価 / 0.65').font = FORMULA_FONT
ws1.cell(row=9, column=1, value='※最高単価 = サロン1-5個価格(定価の65%)').font = Font(name='Arial', size=9, color='666666')

ws1.cell(row=11, column=1, value='判定の優先順位').font = SECTION_FONT
steps = [
    '1. 代理店にKB掛け率が設定されているか? → 未設定なら従来方式(5/13, 1/3)',
    '2. ミカエル or 6+1 を含むか? → ミカエル計算(セット価格/6/0.60)',
    '3. エンジェル or 業務用 を含むか? → 常に65%(数量無関係)',
    '4. 上記以外(店販用単品) → 1-5個:65%  6個以上:60%',
]
for i, s in enumerate(steps):
    ws1.cell(row=12+i, column=1, value=s).font = NORMAL

ws1.cell(row=17, column=1, value='その他の控除項目').font = SECTION_FONT
write_row(ws1, 18, ['項目', '計算方法', ''], HEADER_FONT, HEADER_FILL)
write_row(ws1, 19, ['システム利用料', 'KB対象注文件数 x 300円', ''])
write_row(ws1, 20, ['決済手数料', '対象決済(クレジット/Paid/代金引換)の売上 x 3%', ''], NORMAL, LIGHT_FILL)
write_row(ws1, 21, ['消費税', '(KB合計 - システム利用料 - 決済手数料) x 10%', ''])

ws1.cell(row=23, column=1, value='ソースコード').font = SECTION_FONT
write_row(ws1, 24, ['計算関数', 'src/pages/KickbackManage.jsx → calcKbFromProduct()', ''])
write_row(ws1, 25, ['PDF生成', 'src/lib/generateKickbackPdf.js', ''], NORMAL, LIGHT_FILL)
write_row(ws1, 26, ['代理店設定', 'src/pages/Dealers.jsx', ''])

# ===== Sheet 2 =====
ws2 = wb.create_sheet('サロン価格率')
ws2.sheet_properties.tabColor = 'EA580C'
ws2.column_dimensions['A'].width = 25
ws2.column_dimensions['B'].width = 40
ws2.column_dimensions['C'].width = 25
ws2.column_dimensions['D'].width = 35

ws2.cell(row=1, column=1, value='サロン価格率(商品タイプ別)').font = TITLE_FONT
write_row(ws2, 3, ['商品タイプ', '判定条件', 'サロン価格率', '備考'], HEADER_FONT, HEADER_FILL)
write_row(ws2, 4, ['店販用単品 1-5個', '店販用 かつ 数量1-5', '65%', '通常価格'], NORMAL, LIGHT_FILL)
write_row(ws2, 5, ['店販用単品 6個以上', '店販用 かつ 数量6以上', '60%', 'ボリュームディスカウント'])
write_row(ws2, 6, ['業務用単品', '業務用 を含む', '65%(何個でも一律)', '数量に関係なく一律'], NORMAL, LIGHT_FILL)
write_row(ws2, 7, ['エンジェル', 'セット名にエンジェル含む', '65%(何個でも一律)', '業務用と同じ扱い'])
write_row(ws2, 8, ['ミカエル/6+1セット', 'ミカエル or 6+1 含む', '60%(常に6個計算)', 'セット価格/6で単品逆算'], NORMAL, LIGHT_FILL)

ws2.cell(row=10, column=1, value='代理店原価率(一律・代理店ごと設定)').font = TITLE_FONT
write_row(ws2, 12, ['代理店タイプ', '原価率', '備考', '設定場所'], HEADER_FONT, HEADER_FILL)
write_row(ws2, 13, ['タイプA', '定価 x 45%', '数量に関係なく一律', '代理店管理→詳細→KB掛け率'], NORMAL, LIGHT_FILL)
write_row(ws2, 14, ['タイプB', '定価 x 40%', '数量に関係なく一律', '同上'])

# ===== Sheet 3 =====
ws3 = wb.create_sheet('KB早見表')
ws3.sheet_properties.tabColor = '16A34A'
ws3.column_dimensions['A'].width = 25
ws3.column_dimensions['B'].width = 18
ws3.column_dimensions['C'].width = 18
ws3.column_dimensions['D'].width = 35

ws3.cell(row=1, column=1, value='KB早見表(定価10,000円の商品)').font = TITLE_FONT

ws3.cell(row=3, column=1, value='代理店原価率 45%').font = SECTION_FONT
write_row(ws3, 4, ['商品タイプ', '1個あたりKB', 'セットKB', '計算式'], HEADER_FONT, HEADER_FILL)
write_row(ws3, 5, ['店販単品 1-5個', '=10000*(0.65-0.45)', '', '10,000 x (65%-45%)'], NORMAL, LIGHT_FILL)
write_row(ws3, 6, ['店販単品 6個以上', '=10000*(0.60-0.45)', '', '10,000 x (60%-45%)'])
write_row(ws3, 7, ['業務用/エンジェル', '=10000*(0.65-0.45)', '', '10,000 x (65%-45%)'], NORMAL, LIGHT_FILL)
write_row(ws3, 8, ['ミカエル 1セット', '', '=10000*(0.60-0.45)*6', '10,000 x (60%-45%) x 6'])
for r in range(5, 9):
    ws3.cell(row=r, column=2).number_format = '#,##0"円"'
    ws3.cell(row=r, column=3).number_format = '#,##0"円"'

ws3.cell(row=10, column=1, value='代理店原価率 40%').font = SECTION_FONT
write_row(ws3, 11, ['商品タイプ', '1個あたりKB', 'セットKB', '計算式'], HEADER_FONT, HEADER_FILL)
write_row(ws3, 12, ['店販単品 1-5個', '=10000*(0.65-0.40)', '', '10,000 x (65%-40%)'], NORMAL, LIGHT_FILL)
write_row(ws3, 13, ['店販単品 6個以上', '=10000*(0.60-0.40)', '', '10,000 x (60%-40%)'])
write_row(ws3, 14, ['業務用/エンジェル', '=10000*(0.65-0.40)', '', '10,000 x (65%-40%)'], NORMAL, LIGHT_FILL)
write_row(ws3, 15, ['ミカエル 1セット', '', '=10000*(0.60-0.40)*6', '10,000 x (60%-40%) x 6'])
for r in range(12, 16):
    ws3.cell(row=r, column=2).number_format = '#,##0"円"'
    ws3.cell(row=r, column=3).number_format = '#,##0"円"'

# ===== Sheet 4 =====
ws4 = wb.create_sheet('計算例')
ws4.sheet_properties.tabColor = 'DC2626'
widths = [35, 15, 10, 15, 15, 22, 15, 40]
for i, w in enumerate(widths, 1):
    ws4.column_dimensions[get_column_letter(i)].width = w

ws4.cell(row=1, column=1, value='KB計算例(代理店原価率 45%)').font = TITLE_FONT
headers = ['商品名', 'Bカート単価', '数量', '最高単価', '定価(逆算)', 'サロン価格率', 'KB金額', '計算過程']
write_row(ws4, 3, headers, HEADER_FONT, HEADER_FILL)

# 店販1-5個
write_row(ws4, 4, [
    'ビタミンクリーム50g(店販用単品)', 6500, 3, 6500,
    '=D4/0.65', '65%', '=E4*(0.65-0.45)*C4',
    '定価10,000 x 20% x 3個 = 6,000円'
], NORMAL, GREEN_FILL)

# 店販6個以上
write_row(ws4, 5, [
    'ビタミンクリーム50g(店販用単品)', 6000, 6, 6500,
    '=D5/0.65', '60%', '=E5*(0.60-0.45)*C5',
    '定価10,000 x 15% x 6個 = 9,000円'
])

# 業務用8個
write_row(ws4, 6, [
    'ビタミンクリーム50g(業務用単品)', 6500, 8, 6500,
    '=D6/0.65', '65%(業務用一律)', '=E6*(0.65-0.45)*C6',
    '定価10,000 x 20% x 8個 = 16,000円'
], NORMAL, ORANGE_FILL)

# ミカエル
write_row(ws4, 7, [
    '6+1ビタミンクリーム50g(ミカエル)', 36000, 1, '',
    '=B7/6/0.60', '60%(ミカエル)', '=E7*(0.60-0.45)*6*C7',
    'セット36,000/6/0.60=定価10,000 x 15% x 6 = 9,000円'
], NORMAL, YELLOW_FILL)

# エンジェル6個
write_row(ws4, 8, [
    'ビタミンクリーム50g(エンジェル)', 6500, 6, 6500,
    '=D8/0.65', '65%(エンジェル一律)', '=E8*(0.65-0.45)*C8',
    '定価10,000 x 20% x 6個 = 12,000円'
], NORMAL, GREEN_FILL)

for r in range(4, 9):
    ws4.cell(row=r, column=2).number_format = '#,##0"円"'
    ws4.cell(row=r, column=4).number_format = '#,##0"円"'
    ws4.cell(row=r, column=5).number_format = '#,##0"円"'
    ws4.cell(row=r, column=7).number_format = '#,##0"円"'

# ===== Sheet 5 =====
ws5 = wb.create_sheet('代理店設定項目')
ws5.sheet_properties.tabColor = '9333EA'
ws5.column_dimensions['A'].width = 25
ws5.column_dimensions['B'].width = 45
ws5.column_dimensions['C'].width = 30

ws5.cell(row=1, column=1, value='代理店ごとの設定項目').font = TITLE_FONT
write_row(ws5, 3, ['設定項目', '説明', '設定場所'], HEADER_FONT, HEADER_FILL)
settings = [
    ['KB掛け率(%)', '代理店原価率(45 or 40)。未設定なら従来方式', '代理店管理→詳細'],
    ['自社注文を含める', 'ONなら代理店自身の注文をKB清算に含める(売掛)', '代理店管理→詳細'],
    ['固定調整項目', '業務委託費など毎月の固定加減算', '代理店管理→詳細'],
    ['KBグループ', 'A/B/C グループ分け', '代理店管理→詳細'],
    ['請求書メール', '請求書送付先メールアドレス', '代理店管理→詳細'],
    ['振込先口座', 'KB精算の振込先(銀行名/口座番号等)', '代理店管理→詳細'],
]
for i, s in enumerate(settings):
    write_row(ws5, 4+i, s, NORMAL, LIGHT_FILL if i % 2 == 0 else None)

out = r'C:\Users\takay\OneDrive\Desktop\bomber-admin\docs\KB計算ルール.xlsx'
wb.save(out)
print(f'Saved: {out}')
